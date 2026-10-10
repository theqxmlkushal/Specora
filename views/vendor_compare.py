"""Vendor Compare: heatmap matrix, deviation register, clarifications loop."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
import plotly.graph_objects as go
import streamlit as st

import config
import db
import state as S
import styles

COLOR = {"Compliant": "#27AE60", "Deviating": "#EB5757",
         "Unclear": "#FFB020", "Missing": "#7B61FF"}
CODE = {"Missing": 0, "Deviating": 1, "Unclear": 2, "Compliant": 3}


def render():
    if not st.session_state.get("results"):
        st.info("Load the demo project and extract a PDS first — the approved MDS is "
                "compared against every offer.")
        return
    t1, t2, t3 = st.tabs(["1 · Matrix", "2 · Deviation register", "3 · Clarifications"])
    with t1:
        _matrix()
    with t2:
        _register()
    with t3:
        _clarify()


def _reqs():
    return S.rules_requirements()


def _matrix():
    results = st.session_state["results"]
    reqs = _reqs()
    vendors = sorted({r.vendor for r in results})
    lookup = {(r.vendor, r.req_id): r for r in results}
    # reversed so R01 sits at the top of the chart
    z, txt, y = [], [], []
    for req in reversed(reqs):
        y.append(f"{req['id']} {req['key']}")
        z.append([CODE[lookup[(v, req['id'])].verdict] for v in vendors])
        txt.append([lookup[(v, req['id'])].verdict for v in vendors])
    fig = go.Figure(go.Heatmap(
        z=z, x=vendors, y=y, text=txt, texttemplate="%{text}",
        colorscale=[[0, "#7B61FF"], [0.249, "#7B61FF"],
                    [0.25, "#EB5757"], [0.499, "#EB5757"],
                    [0.5, "#FFB020"], [0.749, "#FFB020"],
                    [0.75, "#27AE60"], [1, "#27AE60"]],
        zmin=-0.5, zmax=3.5,
        colorbar=dict(title="verdict", tickvals=[0, 1, 2, 3],
                      ticktext=["Missing", "Deviating", "Unclear", "Compliant"]),
        hovertemplate="%{y}<br>%{x}: %{text}<extra></extra>"))
    fig.update_layout(height=430, margin=dict(l=10, r=10, t=10, b=10),
                      paper_bgcolor="rgba(0,0,0,0)", font_color="#E6EDF7",
                      xaxis_title="", yaxis_title="")
    st.plotly_chart(fig, use_container_width=True)
    st.caption(f"Requirement sources: " +
               ", ".join(sorted({r['source'] for r in reqs})) +
               " · click-free detail below (select any cell).")

    c1, c2 = st.columns(2)
    vend = c1.selectbox("Vendor", vendors, key="mx_v")
    rid = c2.selectbox("Requirement", [r["id"] for r in reqs], key="mx_r")
    r = lookup[(vend, rid)]
    st.markdown(
        f"<div class='glass'><b>{r.req_id} · {r.key}</b> {styles.chip('GREEN' if r.verdict=='Compliant' else ('RED' if r.verdict in ('Deviating','Missing') else 'AMBER'), r.verdict)}"
        f"<br><span class='sub'>source: <b>{r.source}</b> · category: {r.category} · "
        f"critical: {'yes' if r.critical else 'no'} · weight {r.weight}</span>"
        f"<br>required <code>{r.required}</code> → offered <code>{r.offered}</code>"
        f"<br><span style='font-size:.88rem'>{r.explanation}</span>"
        + (f"<br><span style='font-size:.86rem;color:#FFB020'>📌 {'; '.join(r.notes)}</span>"
           if r.notes else "")
        + (f"<br>❓ <i>{r.question}</i>" if r.question else "")
        + "</div>", unsafe_allow_html=True)


def _register():
    results = st.session_state["results"]
    vendors = sorted({r.vendor for r in results})
    f1, f2, f3 = st.columns(3)
    fv = f1.multiselect("Vendor", vendors, default=vendors, key="dr_v")
    fc = f2.multiselect("Class", ["Compliant", "Deviating", "Unclear", "Missing"],
                        default=["Deviating", "Unclear", "Missing"], key="dr_c")
    fcrit = f3.checkbox("Critical only", key="dr_crit")
    rows = [r for r in results if r.vendor in fv and r.verdict in fc
            and (r.critical or not fcrit)]
    df = pd.DataFrame([{
        "ID": r.req_id, "vendor": r.vendor, "requirement": r.key, "source": r.source,
        "required": r.required, "offered": r.offered, "verdict": r.verdict,
        "critical": "Yes" if r.critical else "No", "category": r.category,
        "weight": r.weight, "proposed resolution": r.question or r.explanation,
    } for r in rows])
    html = df.to_html(escape=True, index=False) if len(df) else "<i>No rows.</i>"
    for k, c in COLOR.items():
        html = html.replace(f">{k}<", f"><span style='color:{c};font-weight:700'>{k}</span><")
    st.markdown('<div style="max-height:340px;overflow:auto;border:1px solid rgba(255,255,255,.12);'
                'border-radius:12px;background:rgba(0,0,0,.22);padding:.3rem">' + html + "</div>",
                unsafe_allow_html=True)
    if st.button("⬇ Export deviation register to Excel"):
        out = config.DATA_DIR / "uploads" / "deviation_register.xlsx"
        out.parent.mkdir(parents=True, exist_ok=True)
        _export_register(df, out)
        with open(out, "rb") as fh:
            st.download_button("Download deviation register", fh, file_name=out.name,
                               mime="application/vnd.openxmlformats-officedocument"
                                    ".spreadsheetml.sheet")
        db.audit(S.get_conn(), "engineer", "export_register", {"rows": len(df)})
        st.toast("Deviation register exported", icon="⬇️")


def _export_register(df: pd.DataFrame, out: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Deviations"
    ws.append(["Specora — deviation register (prototype)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(list(df.columns))
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3B63")
    for row in df.itertuples(index=False):
        ws.append(list(row))
    for i, col in enumerate(df.columns, start=1):
        ws.column_dimensions[ws.cell(row=3, column=i).column_letter].width = min(
            40, max(12, int(df[col].astype(str).str.len().max() if len(df) else 12) + 2))
    wb.save(str(out))


def _clarify():
    results = st.session_state["results"]
    open_items = [r for r in results if r.verdict in ("Unclear", "Missing")]
    st.markdown(f"#### Clarification questions · {len(open_items)} open")
    if not open_items:
        st.success("No open clarifications.")
        return
    conn = S.get_conn()
    replies = db.get_replies(conn)
    for vendor in sorted({r.vendor for r in open_items}):
        with st.expander(f"{vendor} · {len([r for r in open_items if r.vendor == vendor])} "
                         f"questions", expanded=(vendor == "Vendor C")):
            for r in [x for x in open_items if x.vendor == vendor]:
                answered = replies.get((vendor, r.req_id))
                st.markdown(
                    f"<div class='glass' style='padding:.55rem .8rem;margin:.3rem 0'>"
                    f"<b>{r.req_id} {r.key}</b> {styles.chip('AMBER', r.verdict)}"
                    f"<br><span style='font-size:.88rem'>❓ {r.question}</span>"
                    + (f"<br><span style='font-size:.86rem;color:#6FCF97'>↩ {answered}</span>"
                       if answered else "")
                    + "</div>", unsafe_allow_html=True)
                if r.question and not answered:
                    with st.form(key=f"reply_{vendor}_{r.req_id}"):
                        ans = st.text_input("Simulate vendor reply", key=f"ra_{r.req_id}")
                        if st.form_submit_button("Send reply & re-evaluate",
                                                 use_container_width=True):
                            if ans.strip():
                                db.save_reply(conn, vendor, r.req_id, ans.strip())
                                db.audit(conn, vendor, "vendor_reply",
                                         {"req": r.req_id, "answer": ans.strip()})
                                S.run_compare()
                                st.toast(f"{vendor} replied to {r.req_id}", icon="📨")
                                st.rerun()
    if replies and st.button("Clear all simulated replies"):
        db.clear_replies(conn)
        db.audit(conn, "engineer", "clear_replies", {})
        S.run_compare()
        st.rerun()
